"""
Recursive chunking of an Excel (.xlsx) file with LangChain.

Install:  pip install langchain-text-splitters openpyxl
Run:      python recursive_chunking_xlsx.py
"""
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook

XLSX_PATH = "posh_applicability_checker_expanded.xlsx"


def get_xlsx_chunks(path=XLSX_PATH):
    # 1. Load the workbook (one Document per sheet, with sheet name in metadata).
    #    Each row becomes one line of text, with its cells joined by " | ".
    workbook = load_workbook(path, data_only=True)  # data_only=True -> formula results, not formulas
    sheets = []
    for ws in workbook.worksheets:
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:  # skip empty rows
                lines.append(" | ".join(cells))
        sheets.append(Document(page_content="\n".join(lines), metadata={"source": path, "sheet": ws.title}))

    # 2. Recursive splitter: tries to split between rows first, then between cells,
    #    then words, then characters, until every chunk fits within chunk_size.
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,       # max characters per chunk
        chunk_overlap=50,     # characters shared between neighbouring chunks
        separators=["\n", " | ", " ", ""],
    )
    return splitter.split_documents(sheets)


if __name__ == "__main__":
    # 3. Display the chunks
    chunks = get_xlsx_chunks()
    print(f"Created {len(chunks)} chunks from {XLSX_PATH}\n")
    for i, chunk in enumerate(chunks, start=1):
        sheet = chunk.metadata["sheet"]
        print(f"===== Chunk {i}  (sheet '{sheet}', {len(chunk.page_content)} chars) =====")
        print(chunk.page_content)
        print()
